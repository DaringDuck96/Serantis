"""Turn the Monster Batch 3.0 Book (xlsx) into Ahrivaea creature definition files.

Usage:  python3 tools/build_batch3.py <book.xlsx> <output creature folder>

The book supplies name, influence (element), rarity, biomes and drops. Everything the
book doesn't cover (stats, movement, size, jobs, levels, spawn weight) is set here so it
follows the same rules as the original 239 creatures:
  * level range and spawn weight are fixed per rarity (copied from the existing creatures)
  * stat totals stay within 500-600 and no single stat goes above 100
  * job descriptions reuse the mod's existing wording
"""
import json, os, re, sys, zlib
import openpyxl

# ---- rules copied from the existing creatures -------------------------------------------
LEVELS = {"Common": (2, 12), "Uncommon": (6, 18), "Rare": (12, 26), "Super Rare": (20, 35), "Legendary": (30, 45)}
WEIGHT = {"Common": 400, "Uncommon": 180, "Rare": 80, "Super Rare": 30, "Legendary": 10}
TOTAL = {"Common": (515, 535), "Uncommon": (525, 545), "Rare": (540, 560), "Legendary": (580, 595)}
SIZE = {"tiny": 1.2, "small": 1.6, "medium": 2.4, "large": 4.2, "huge": 7.5, "colossal": 11.0, "dragon": 12.0}
TOP_SIZE = {"tiny": 0.7, "small": 0.7, "medium": 1.0, "large": 1.0, "huge": 1.6, "colossal": 1.6, "dragon": 1.6}
JOB_TEXT = {
    "defense": "Base defense", "ranching": "Hostile-mob drop collection", "fishing": "Fishing",
    "mining": "Mining stone and ores", "logging": "Tree farming", "brewing": "Brewing potions",
    "enchanting": "Enchanting support", "ice": "Snow and ice harvesting", "farming": "Crop harvesting",
    "growth": "Crop growth support", "redstone": "Redstone device operation", "smelting": "Smelting ores",
    "transport": "Item pickup", "crafting": "Crafting-table assistance",
    "lighting": "Dark-area exploration and lighting", "riding": "Mounted travel support",
}
STATS = ["hp", "attack", "defense", "speed", "magic", "luck", "charisma"]
SHAPES = {  # relative weight of each stat:  hp  atk  def  spd  mag  lck  cha
    "brute":     (1.20, 1.35, 1.15, 0.80, 0.60, 0.90, 0.80),
    "tank":      (1.30, 0.95, 1.45, 0.60, 0.80, 0.90, 0.80),
    "mage":      (0.85, 0.60, 0.80, 1.00, 1.50, 1.05, 1.10),
    "striker":   (0.85, 1.25, 0.75, 1.45, 0.80, 1.00, 0.90),
    "trickster": (0.80, 0.80, 0.75, 1.25, 1.10, 1.30, 1.00),
    "dancer":    (0.90, 0.70, 0.80, 1.10, 1.15, 1.05, 1.40),
    "shadow":    (0.90, 1.10, 0.85, 1.20, 1.00, 1.00, 0.80),
    "beast":     (1.10, 1.20, 1.00, 1.10, 0.70, 0.95, 0.95),
    "balanced":  (1, 1, 1, 1, 1, 1, 1),
    "dragon":    (1.10, 1.15, 1.00, 1.00, 1.05, 0.85, 0.90),
}

# ---- per-creature design: movement, size, stat shape, two base jobs ----------------------
D = {
 # Fire
 "Ashback Devourer": ("walk","huge","brute","smelting","defense"),
 "Caldera Wyrm": ("walk","large","beast","smelting","mining"),
 "Cinderwing": ("fly","small","striker","smelting","transport"),
 "Ember Magus": ("walk","medium","mage","brewing","enchanting"),
 "Magma Crawler": ("walk","small","beast","smelting","mining"),
 "Cinder Shadow": ("walk","medium","shadow","defense","ranching"),
 "Magma Bulwark": ("walk","large","tank","defense","smelting"),
 "Blaze Dancer": ("walk","medium","dancer","lighting","brewing"),
 # Arcane
 "Aether Beamer": ("fly","medium","mage","redstone","enchanting"),
 "Ancient Spell Engine": ("walk","large","mage","redstone","crafting"),
 "Glyph Guardian": ("walk","medium","tank","defense","enchanting"),
 "Rune Magus": ("walk","medium","mage","enchanting","brewing"),
 "Runic Colossus": ("walk","colossal","tank","defense","mining"),
 "Rune Crawler": ("walk","small","trickster","redstone","transport"),
 # Nature
 "Bloom Eye Horror": ("walk","large","mage","growth","farming"),
 "Canopy Skullspider": ("walk","medium","striker","ranching","defense"),
 "Elderwood Guardian": ("walk","large","tank","logging","defense"),
 "Snapvine": ("walk","small","beast","farming","growth"),
 "Thorncoil": ("walk","medium","beast","growth","defense"),
 "Thorn Shadow": ("walk","medium","shadow","defense","logging"),
 # Water
 "Abyssal Maw": ("swim","large","brute","fishing","defense"),
 "Reef Spitter": ("swim","small","striker","fishing","transport"),
 "River Zorling": ("swim","small","trickster","fishing","transport"),
 "Tidecore Horror": ("swim","huge","mage","fishing","defense"),
 "Whirlpool Leech": ("swim","medium","beast","fishing","ranching"),
 "Tide Shadow": ("swim","medium","shadow","fishing","defense"),
 "Coral Dancer": ("swim","medium","dancer","fishing","growth"),
 "Bubble Stalker": ("swim","medium","striker","fishing","defense"),
 "Drill Shark": ("swim","large","brute","fishing","mining"),
 # Darkness
 "Dread Knight": ("walk","medium","brute","defense","ranching"),
 "Dreadwalker": ("walk","medium","shadow","defense","ranching"),
 "Gloomwing": ("fly","large","striker","defense","transport"),
 "Gravewrap": ("walk","medium","tank","ranching","defense"),
 "Umbral Revenant": ("walk","medium","mage","ranching","enchanting"),
 "Nightmare Stalker": ("walk","medium","striker","defense","ranching"),
 "Gloom Shadow": ("walk","medium","shadow","defense","ranching"),
 "Grave Lurker": ("walk","large","brute","ranching","mining"),
 # Fae
 "Glimmer Sprite": ("fly","tiny","trickster","growth","lighting"),
 "Lantern Trickster": ("fly","small","trickster","lighting","transport"),
 "Masked Puck": ("walk","small","trickster","brewing","transport"),
 "Moonlit Stag": ("walk","large","beast","transport","growth"),
 "Moonwing": ("fly","small","mage","growth","lighting"),
 "Dream Dancer": ("walk","medium","dancer","brewing","growth"),
 "Phantom Jester": ("walk","medium","trickster","brewing","ranching"),
 # Radiant
 "Crowned Gleeok": ("fly","huge","mage","defense","smelting"),
 "Dawn Magus": ("walk","medium","mage","enchanting","brewing"),
 "Lumina Wyrm": ("walk","large","beast","lighting","defense"),
 "Sanctum Sentinel": ("walk","large","tank","defense","enchanting"),
 "Sunwatcher": ("fly","medium","mage","growth","lighting"),
 "Celestial Watcher": ("fly","large","mage","enchanting","lighting"),
 "Sunlit Dancer": ("walk","medium","dancer","growth","farming"),
 # Earth
 "Basalt Maw": ("walk","large","brute","mining","smelting"),
 "Burrowcoil": ("walk","medium","beast","mining","farming"),
 "Oreback Titan": ("walk","colossal","tank","mining","defense"),
 "Prism Sentinel": ("walk","large","tank","mining","enchanting"),
 "Ruin Sentinel": ("walk","large","tank","defense","crafting"),
 "Stone Shadow": ("walk","medium","shadow","mining","defense"),
 "Quake Brute": ("walk","large","brute","mining","defense"),
 "Burrowing Horror": ("walk","large","beast","mining","ranching"),
 # None
 "Gear Guzzler": ("walk","medium","tank","redstone","crafting"),
 "Ironjaw Brute": ("walk","large","brute","mining","defense"),
 "One Eye Colossus": ("walk","colossal","brute","defense","mining"),
 "Pebble Spitter": ("walk","small","striker","mining","transport"),
 "Scrap Goblin": ("walk","small","trickster","crafting","transport"),
 "Hollow Grunt": ("walk","medium","balanced","defense","logging"),
 "Crawling Mimic": ("walk","medium","trickster","transport","ranching"),
 "Ironclad Ravager": ("walk","large","brute","defense","transport"),
 "One-Eyed Devourer": ("walk","huge","brute","ranching","defense"),
 "Pale Duelist": ("walk","medium","striker","defense","crafting"),
 # Frost
 "Glacier Spitter": ("walk","medium","striker","ice","defense"),
 "Rime Magus": ("walk","medium","mage","ice","brewing"),
 "Rimewing": ("fly","small","striker","ice","transport"),
 "Shardheart": ("walk","large","tank","ice","mining"),
 "Snowfang": ("walk","medium","beast","ice","defense"),
 "Rime Dancer": ("walk","medium","dancer","ice","growth"),
 # Storm
 "Stormblade Phantom": ("fly","medium","striker","defense","redstone"),
 "Stormjaw Leviathan": ("swim","huge","brute","fishing","redstone"),
 "Tempest Roc": ("fly","huge","striker","transport","defense"),
 "Thunder Magus": ("walk","medium","mage","redstone","enchanting"),
 "Voltwing": ("fly","small","striker","redstone","transport"),
 "Tempest Flyer": ("fly","medium","striker","transport","redstone"),
 # Void
 "Dimensional Duelist": ("walk","medium","striker","defense","mining"),
 "Hollow Phantom": ("fly","medium","shadow","ranching","defense"),
 "Rift Clutch": ("walk","large","mage","mining","enchanting"),
 "Rift Wraith": ("fly","medium","mage","ranching","transport"),
 "Void Hand": ("walk","medium","trickster","transport","mining"),
 "Rift Shadow": ("walk","medium","shadow","defense","logging"),
 # Summoned bosses
 "Harbinger of Calamity": ("walk","colossal","brute","defense","ranching"),
 "False Hydra": ("walk","colossal","mage","enchanting","defense"),
}
DRAGON_JOBS = {"none": ("transport","defense"), "fire": ("smelting","defense"), "frost": ("ice","defense"),
    "storm": ("redstone","transport"), "earth": ("mining","defense"), "water": ("fishing","transport"),
    "nature": ("growth","logging"), "void": ("mining","transport"), "fae": ("growth","lighting"),
    "darkness": ("ranching","defense"), "arcane": ("enchanting","brewing"), "radiant": ("lighting","enchanting")}
DRAGON_MOVE = {"Mokele mbembe": "walk", "Ladon": "walk"}           # land/river dragons
DRAGON_FOCUS = {"Rendmaw": "attack", "Arcanum": "magic", "Bastion": "defense",
                "Galewing": "speed", "Fortune": "luck", "Regalia": "charisma"}
APEX = {"The First Dragon", "The Last Dragon"}                     # highest stats in the game

# ---- book text -> Minecraft IDs ----------------------------------------------------------
BIOMES = set("""badlands bamboo_jungle basalt_deltas beach birch_forest cherry_grove cold_ocean crimson_forest
dappled_forest dark_forest deep_cold_ocean deep_dark deep_frozen_ocean deep_lukewarm_ocean deep_ocean desert
dripstone_caves end_barrens end_highlands end_midlands eroded_badlands flower_forest forest frozen_ocean
frozen_peaks frozen_river grove ice_spikes jagged_peaks jungle lukewarm_ocean lush_caves mangrove_swamp meadow
mushroom_fields nether_wastes ocean old_growth_birch_forest old_growth_pine_taiga old_growth_spruce_taiga
pale_garden plains river savanna savanna_plateau small_end_islands snowy_beach snowy_plains snowy_slopes
snowy_taiga soul_sand_valley sparse_jungle stony_peaks stony_shore sunflower_plains swamp taiga the_end
the_void warm_ocean warped_forest windswept_forest windswept_gravelly_hills windswept_hills windswept_savanna
wooded_badlands""".split())
BIOME_FIX = {"crimson_forests": "crimson_forest", "dark_oak_forest": "dark_forest",
             "old_growth_taiga": "old_growth_spruce_taiga"}
NOT_BIOMES = {"ancient_city", "amethyst_geode"}  # structures/features, not biomes; their regions are covered by the other biomes listed
ITEMS = {"Raw Cod": "cod", "Prismarine Crystal": "prismarine_crystals"}

NAME_FIXES = {"Scrap Gobblin": "Scrap Goblin"}   # spelling corrections to the book

def cid(name): return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")

def biomes(text, log):
    out = []
    for b in text.split(","):
        b = BIOME_FIX.get(cid(b.replace("&#x20;", " ")), cid(b.replace("&#x20;", " ")))
        if b in NOT_BIOMES: log.append(f"dropped '{b}' (not a biome)"); continue
        if b not in BIOMES: raise SystemExit(f"Unknown biome: {b}")
        if "minecraft:" + b not in out: out.append("minecraft:" + b)
    return out

def drops(text):
    out = []
    for part in filter(None, (p.strip() for p in (text or "").split(";"))):
        m = re.fullmatch(r"(.+?)\s*\((\d+)%\s*\|\s*(\d+)(?:-(\d+))?\)", part)
        if not m: raise SystemExit(f"Can't read drop: {part}")
        name, pct, lo, hi = m.groups()
        out.append({"item": "minecraft:" + ITEMS.get(name, cid(name)), "chance": int(pct) / 100,
                    "min": int(lo), "max": int(hi or lo)})
    return out

def make_stats(key, shape, total, focus=None):
    w = list(SHAPES[shape])
    seed = zlib.crc32(key.encode())
    w = [x * (0.94 + ((seed >> (i * 4)) & 15) / 125) for i, x in enumerate(w)]   # small fixed per-creature variation
    if focus: w[STATS.index(focus)] = max(w) * 2.2
    s = [max(35, min(100, round(x / sum(w) * total))) for x in w]
    order = sorted(range(7), key=lambda i: -w[i])
    while sum(s) != total:                       # nudge to the exact total within 35..100
        step = 1 if sum(s) < total else -1
        for i in (order if step > 0 else order[::-1]):
            if 35 <= s[i] + step <= 100 and sum(s) != total: s[i] += step
    return dict(zip(STATS, s))

def total_for(key, rarity):
    lo, hi = TOTAL[rarity]; return lo + zlib.crc32(key.encode()) % (hi - lo + 1)

def creature(name, influence, rarity, move, size, shape, jobs, spawn, drop_list, focus=None, total=None):
    key = cid(name)
    stats = make_stats(key, shape, total or total_for(key, rarity), focus)
    form = {"stats": stats, "rarity": rarity}
    if spawn: form["spawn"] = spawn
    form.update({"drops": drop_list, "jobs": [{"job": j, "text": JOB_TEXT[j]} for j in jobs],
                 "size": SIZE[size], "wave": 3})
    return key, {"name": name, "forms": {influence: form}, "stats": stats, "size": TOP_SIZE[size], "movement": move}

def main(book, outdir):
    wb = openpyxl.load_workbook(book, data_only=True)
    os.makedirs(outdir, exist_ok=True)
    made, notes = {}, []
    for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        if not row[1]: continue
        infl, raw, rar, bio, drp = row
        name = re.sub(r"\s+", " ", raw).strip()
        name = re.sub(r"\s*\(.*\)$", "", name)                  # "Drill Shark (Abyssal Maw)" -> "Drill Shark"
        name = NAME_FIXES.get(name, name)
        rarity = {"Commmon": "Common", "Absolute Legendary": "Legendary"}.get(rar.strip(), rar.strip())
        move, size, shape, *jobs = D[name]
        log = []
        if rar.strip() == "Absolute Legendary":
            spawn = None; notes.append(f"{name}: summon-only in the book; no natural spawn")
        else:
            lo, hi = LEVELS[rarity]
            spawn = {"biomes": biomes(bio, log), "min_level": lo, "max_level": hi, "weight": WEIGHT[rarity]}
        notes += [f"{name}: {m}" for m in log]
        made[name] = creature(name, infl.lower(), rarity, move, size, shape, jobs, spawn, drops(drp))
    for row in wb.worksheets[1].iter_rows(min_row=2, values_only=True):
        if not row[1]: continue
        infl, raw = row[0].lower(), row[1].strip()
        name = re.sub(r"\s*\(.*\)$", "", raw)
        focus = DRAGON_FOCUS.get(name)
        total = 600 if name in APEX else (595 if focus else None)
        made[name] = creature(name, infl, "Legendary", DRAGON_MOVE.get(name, "fly"),
                              "dragon", "balanced" if name in APEX else "dragon",
                              DRAGON_JOBS[infl], None, drops(row[4]), focus, total)
    for name, (key, data) in made.items():
        with open(os.path.join(outdir, key + ".json"), "w") as f:
            json.dump(data, f, indent=2); f.write("\n")
    print(f"wrote {len(made)} creatures to {outdir}")
    for n in notes: print("  note:", n)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
